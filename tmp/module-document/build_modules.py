from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'output' / 'documents'
OUT.mkdir(parents=True, exist_ok=True)

# Original nine scope module identifiers are retained from mvp.md.
modules = [
('User Management and Access Control', [
'Secure login and logout with password hashing and database-backed sessions.',
'Separate Admin and Agent workspaces; Manager access to administrative functions.',
'Client assignments, selected-client access controls and organization isolation.',
'User, client and integration administration with recorded audit activity.',
], 'Supporting platform capability. Independent agent signup, email verification and Google sign-in remain deferred on a separate branch.'),
('6.1 Multi CRM Sync Engine', [
'HubSpot agency CRM with unified contacts, companies, deals, owners and sales activities.',
'Client-specific Zoho CRM records, synchronization and supported record editing.',
'Closed Won handoff from agency deals to the relevant client CRM with duplicate protection and retry.',
'Provider connections, encrypted credentials, synchronization history and manual Sync or Retry controls.',
], 'Full automatic cross-system synchronization and conflict recovery still require completion and verification. External writes do not form one atomic transaction.'),
('6.2 VoIP and SMS Identity Gateway', [
'Administrator assignment of client-specific Twilio numbers and agent identity selection.',
'Outbound SMS and phone calls that ring the agent first, then bridge the selected contact.',
'Signed incoming SMS and voice webhooks, client routing and optional incoming-call forwarding.',
'Delivery status, call outcomes and duration recorded in the timeline; manual call and SMS links remain available.',
], 'Implemented in the current branch. Live Twilio verification requires configured credentials, numbers and a public callback address.'),
('6.3 Unified Interaction Aggregator', [
'A client-specific timeline linking interactions to contacts, companies and deals.',
'Imported HubSpot calls, meetings, tasks, notes and permitted email activities.',
'Incoming and outgoing Twilio events with status updates and duplicate protection.',
'Manual interaction and note entry, filtering, and refresh of communication activity.',
], 'HubSpot email history depends on provider permission. Dedicated mailbox ingestion and automatic call transcription remain pending.'),
('6.4 AI Intelligence and Lead Analysis', [
'Transcript entry linked to an active-client contact and stored analysis history.',
'Rule-based extraction of budget, timeline, requirements, intent and objections.',
'Explainable lead score, lead temperature and deal-probability estimates.',
'Manual review and correction of analysis fields and scores with audit records.',
], 'The current analysis uses deterministic rules. AI model integration, automatic transcript ingestion and approved extraction-to-CRM updates remain pending.'),
('6.5 Sales Script Architect', [
'Context-based script drafts using client branding, contact details and saved lead analysis.',
'Template generation for call, SMS and email channels with a selected purpose.',
'Saved scripts and reusable templates available in the agent workspace.',
'Manual script creation, editing and override with client-scoped storage.',
], 'Generation is template-based. AI-written and adaptive sales scripts remain pending.'),
('6.6 Commission and Financial Tracker', [
'Automatic commission creation for attributed Closed Won deals with a closing snapshot.',
'Rate selection from deal, agent, client or agency defaults.',
'Expected, received and pending balances with payment entries and audited manual corrections.',
'Agent and agency reporting, inclusive closing-date filters, currency-separated totals and CSV export.',
], 'Live CRM verification remains necessary. Payment-date cash-flow reports and historical balance snapshots are separate future work.'),
('6.7 Sales Ops Reconciliation', [
'Automatic active-work tracking by agent and selected client with Pause and Resume controls.',
'Focus and inactivity handling, server-calculated time and protection against overlapping sessions.',
'Manual time entries, correction of recorded work and reconciliation of operational totals.',
'Personal closed-deal tracker with optional ClickUp task synchronization and retry; financial reconciliation uses commission balances.',
], 'Core time tracking and reconciliation are implemented. Offline work requires manual entry; live ClickUp verification remains provider-dependent.'),
('6.8 Workflow and Smart Scheduling', [
'Client-linked local appointments with create, edit, confirm and cancel actions.',
'Agent-owned Google Calendar, Outlook and Calendly OAuth connections.',
'Supported external booking, update or cancellation, optional invitations and synchronization retry or refresh.',
'Calendly individual event-type management and weekly availability editing; manual scheduling remains available.',
], 'Provider verification is still required, including Outlook. AI meeting-intent detection and extended Calendly question/location handling remain pending.'),
('6.9 Document and Proposal Generator', [
'Template-based proposals, follow-up email drafts and message drafts using client and lead context.',
'Editable stored drafts, manual documents and draft/final status management.',
'Linking documents to contacts, lead analyses and deals within the selected client.',
'PDF upload, storage and download with audit records; uploaded PDFs are replaced by re-uploading revised files.',
], 'Draft generation is template-based. AI proposal drafting remains pending.'),
]

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin, s.bottom_margin = Inches(.55), Inches(.55)
s.left_margin, s.right_margin = Inches(.68), Inches(.68)
s.header_distance, s.footer_distance = Inches(.22), Inches(.22)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet']:
    st = doc.styles[name]
    st.font.name = 'Times New Roman'
    st.font.size = Pt(10.5)
    st.font.color.rgb = RGBColor(0,0,0)
    fonts = st.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        fonts.attrib.pop(qn('w:'+attr), None)
    fonts.set(qn('w:ascii'), 'Times New Roman')
    fonts.set(qn('w:hAnsi'), 'Times New Roman')
    for e in list(st.element.xpath('./w:pPr/w:pBdr')):
        e.getparent().remove(e)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.line_spacing = 1.0
    if name != 'Normal':
        st.paragraph_format.keep_with_next = False
doc.styles['Title'].font.size = Pt(12)
doc.styles['Title'].font.bold = True
doc.core_properties.title = 'Multi CRM Unified Sales Platform Module Document'
doc.core_properties.subject = 'Built modules and remaining deliverables'
doc.core_properties.author = ''

def prevent_split(row):
    trpr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:cantSplit'); trpr.append(e)

def configure_table(t, widths=(2.18,4.96)):
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    pr = t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for k in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+k)
        for a,v in [('val','single'),('sz','6'),('color','000000')]: x.set(qn('w:'+a),v)
        borders.append(x)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for k,n in [('top',65),('bottom',65),('left',90),('right',90)]:
        x=OxmlElement('w:'+k); x.set(qn('w:w'),str(n)); x.set(qn('w:type'),'dxa'); margins.append(x)
    pr.append(margins)
    for row in t.rows:
        prevent_split(row)
        for cell,width in zip(row.cells,widths):
            cell.width=Inches(width)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP

def cell_text(c, text, bold=False, center=False, style=None):
    p=c.paragraphs[0]
    if style: p.style=style
    p.paragraph_format.space_after=Pt(0)
    if center: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(text); r.bold=bold
    r.font.name='Times New Roman'
    r.font.color.rgb=RGBColor(0,0,0)
    if style=='Title':
        borders=OxmlElement('w:pBdr')
        for edge in ['top','left','bottom','right','between','bar']:
            e=OxmlElement('w:'+edge); e.set(qn('w:val'),'nil'); borders.append(e)
        p._p.get_or_add_pPr().append(borders)
    return p

def module_table(selected):
    t=doc.add_table(rows=1,cols=2)
    cell_text(t.cell(0,0),'Module Name',True,True)
    cell_text(t.cell(0,1),'Module Description with Deliverables',True,True)
    e=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(e)
    for name,items,status in selected:
        row=t.add_row(); cell_text(row.cells[0],name,True)
        c=row.cells[1]
        for i,text in enumerate(items):
            p=c.paragraphs[0] if i==0 else c.add_paragraph()
            p.paragraph_format.left_indent=Inches(.14)
            p.paragraph_format.first_line_indent=Inches(-.12)
            p.paragraph_format.space_after=Pt(3)
            p.add_run('\u2022  '+text)
        p=c.add_paragraph(); p.paragraph_format.space_after=Pt(0)
        p.add_run('Status  ').bold=True
        p.add_run(status).italic=True
    configure_table(t)
    return t

intro=doc.add_paragraph()
intro.paragraph_format.space_after=Pt(5)
rr=intro.add_run('Draft prepared on 4 October 2026 from the project MVP specification and current implementation. Original scope document pending confirmation. Built features and remaining deliverables are distinguished below; external integrations are not certified as live-tested by this document.')
rr.font.size=Pt(9)
rr.italic=True

t=doc.add_table(rows=8,cols=2)
cell_text(t.cell(0,0),'Project Title',True,True)
cell_text(t.cell(0,1),'Multi CRM Unified Sales Platform',True,True,'Title')
cell_text(t.cell(1,0),'Supervisor Name',True,True)
cell_text(t.cell(2,0),'Co Supervisor Name',True,True)
cell_text(t.cell(3,0).merge(t.cell(3,1)),'Group Information',True,True)
cell_text(t.cell(4,0),'Roll Number',True,True)
cell_text(t.cell(4,1),'Student Name',True,True)
for row in t.rows[1:]:
    row.height=Inches(.23); row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
configure_table(t)
doc.add_paragraph().paragraph_format.space_after=Pt(0)
module_table(modules[:3])
doc.add_page_break()
module_table(modules[3:7])
doc.add_page_break()
module_table(modules[7:])

p=doc.add_paragraph('Evaluator Feedback')
p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(3)
p.runs[0].bold=True
p=doc.add_paragraph('Please record recommended improvements or additional modules in the section below.')
p.paragraph_format.space_after=Pt(5)
t=doc.add_table(rows=2,cols=2)
cell_text(t.cell(0,0),'Evaluator Name',True)
cell_text(t.cell(0,1),'Suggestions',True,True)
t.rows[1].height=Inches(.42); t.rows[1].height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
configure_table(t)
p=doc.add_paragraph()
p.paragraph_format.space_before=Pt(11)
p.add_run('Supervisor Signature  __________________')
p.add_run('     FYP Committee Member Signature  __________________')
p.runs[0].font.size=Pt(10); p.runs[1].font.size=Pt(10)

p=doc.add_paragraph('Source basis  mvp.md; MVP_PROGRESS.md; current application services; Twilio, calendar, commission reconciliation and active time tracking documentation. The nine module numbers follow the project specification; user management is included as a supporting capability.')
p.paragraph_format.space_before=Pt(5)
p.paragraph_format.space_after=Pt(0)
for r in p.runs: r.font.size=Pt(8)

p=s.footer.paragraphs[0]
p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=p.add_run('Module Document Draft  |  '); r.font.size=Pt(8)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); p._p.append(field)

path=OUT/'Multi_CRM_Module_Document_Draft.docx'
doc.save(path)
print(path)
print('Modules:',len(modules),'Words:',sum(len(p.text.split()) for table in doc.tables for row in table.rows for c in row.cells for p in c.paragraphs))
